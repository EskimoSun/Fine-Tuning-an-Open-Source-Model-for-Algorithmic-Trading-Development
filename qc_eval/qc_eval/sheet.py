"""
Builds and fills the eval spreadsheet (Legend / Results / Summary).

Row order on the Results tab is the order of prompts_100.jsonl, starting
at row 3, so results can be written back by prompt id.
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FIRST_ROW = 3
MODELS = ["fable", "qwen_base", "qwen_tuned"]
# columns per model: compile errors, runtime errors, fidelity
MODEL_COLS = {"fable": ("F", "G", "H"), "qwen_base": ("I", "J", "K"), "qwen_tuned": ("L", "M", "N")}
HUMAN_FLAG_COL, HUMAN_SCORE_COL, NOTES_COL = "O", "P", "Q"
SEGMENT_COL, SPEC_COL = "C", "D"

_F = "Arial"
_thin = Side(style="thin", color="BFBFBF")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
_wrap = Alignment(vertical="top", wrap_text=True)
_hdr = PatternFill("solid", fgColor="1F3A5F")
_input = PatternFill("solid", fgColor="FFFF00")
_group = {"Fable": "D9E2F3", "Qwen base": "E2EFDA", "Qwen tuned": "FFF2CC", "Human": "FCE4D6"}


def _font(**k):
    return Font(name=_F, size=k.pop("size", 10), **k)


def _legend(wb):
    lg = wb.active
    lg.title = "Legend"
    rows = [
        ("QuantConnect Fine-Tuning Eval: Prompt x Error Sheet", None),
        ("", None),
        ("Purpose", "Count errors per prompt for Fable (reference), Qwen before fine-tuning and Qwen after fine-tuning, to show error reduction."),
        ("Layout", "Results tab: one row per prompt. 10 code templates x 10 prompt versions = 100 rows. Column A is the template file (QuantConnect project-templates repo), included under data/templates/."),
        ("Prompt grid", "Versions 1-9 = 3 customer segments x 3 specificity levels. Version 10 = ceiling prompt (near-pseudocode description of the template; every model should score well on it)."),
        ("Segments", "individual: hobbyist vocabulary, single ticker, no constraints. startup: terse researcher shorthand, in a hurry. institutional: spec-style, mentions risk limits, compliance, matching an internal implementation."),
        ("Specificity", "vague: key parameters left out, model must choose sensibly. precise: every parameter needed to reproduce the template. over_specified: precise plus extra requirements that do not change trading (logging, plotting, naming, style)."),
        ("How it's filled", "pipeline.py writes every score column automatically. Only the yellow Human columns are filled by hand."),
        ("Compile errors", "Number of error lines in the LEAN build log from the QC compile API. 0 if it builds."),
        ("Runtime errors", "1 if the QC backtest crashed, 0 if it completed. A crash stops the backtest, so this is 0 or 1. Blank if it did not compile."),
        ("Fidelity to prompt (1-5)", "Scored by Fable as judge, temperature 0, using this rubric:"),
        ("   5", "Every requested asset, indicator, parameter, entry/exit rule and date range is present and correct."),
        ("   4", "All key elements present; one minor detail off (e.g. wrong rebalance time)."),
        ("   3", "Core idea right but one key element wrong or missing (e.g. wrong SMA length, wrong exit rule)."),
        ("   2", "Recognizably the same strategy family, but several key elements wrong."),
        ("   1", "Does not implement the requested strategy."),
        ("Vague prompts", "Vague versions leave details out on purpose. Judge them on whether the code is a sensible version of the request, not on missing parameters."),
        ("Human check", "pipeline.py marks 10 random rows Y. A team member reads Qwen base's code for those rows and enters a 1-5 fidelity score in column P. Summary compares it to Fable's score."),
        ("Notes", "Short error detail written by the pipeline (first compile error, runtime error, static-check warnings)."),
        ("", None),
        ("Data rule", "These 100 prompts are TEST ONLY. Do not use them, or Fable's reference strategies, as training data."),
        ("Reproducibility", "All generations use temperature 0 and are cached in runs/. Each cached record stores the model id and timestamp, so reruns score the same outputs unless you pass --regenerate."),
    ]
    for i, (a, b) in enumerate(rows, 1):
        lg.cell(i, 1, a).font = _font(bold=bool(b) or i == 1, size=14 if i == 1 else 10)
        if b:
            lg.cell(i, 2, b).font = _font()
            lg.cell(i, 2).alignment = _wrap
    lg.column_dimensions["A"].width = 26
    lg.column_dimensions["B"].width = 100


def _results(wb, prompts):
    ws = wb.create_sheet("Results")
    groups = [("Prompt", 1, 5), ("Fable (reference)", 6, 8), ("Qwen base (pre fine-tune)", 9, 11),
              ("Qwen tuned (post fine-tune)", 12, 14), ("Human check", 15, 16), ("", 17, 17)]
    for name, c1, c2 in groups:
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
        cell = ws.cell(1, c1, name)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.alignment = _center
        for c in range(c1, c2 + 1):
            ws.cell(1, c).fill = _hdr
    headers = ["Code template", "Version", "Segment", "Specificity", "Prompt text"] + \
              ["Compile errors", "Runtime errors", "Fidelity (1-5)"] * 3 + \
              ["Human checked (Y/N)", "Human fidelity, Qwen base (1-5)", "Notes"]
    keys = [None] * 5 + ["Fable"] * 3 + ["Qwen base"] * 3 + ["Qwen tuned"] * 3 + ["Human"] * 2 + [None]
    for c, (h, k) in enumerate(zip(headers, keys), 1):
        cell = ws.cell(2, c, h)
        cell.font = _font(bold=True)
        cell.alignment = _center
        cell.border = _border
        cell.fill = PatternFill("solid", fgColor=_group[k] if k else "D9D9D9")

    last = FIRST_ROW + len(prompts) - 1
    for i, p in enumerate(prompts):
        r = FIRST_ROW + i
        tpl = p.get("reference_path", p["template"]).replace("data/templates/", "")
        for c, v in enumerate([tpl, p["version"], p.get("segment", ""), p.get("specificity", ""), p["prompt"]], 1):
            ws.cell(r, c, v)
        for c in range(1, 18):
            cell = ws.cell(r, c)
            cell.font = _font()
            cell.border = _border
            cell.alignment = _wrap if c in (5, 17) else _center
        for col in (HUMAN_FLAG_COL, HUMAN_SCORE_COL):
            ws[f"{col}{r}"].fill = _input
        if (i // 10) % 2:
            for c in (1, 2, 3, 4):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F2F2F2")

    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_fid = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv_yn)
    ws.add_data_validation(dv_fid)
    dv_yn.add(f"{HUMAN_FLAG_COL}{FIRST_ROW}:{HUMAN_FLAG_COL}{last}")
    dv_fid.add(f"{HUMAN_SCORE_COL}{FIRST_ROW}:{HUMAN_SCORE_COL}{last}")

    widths = [38, 8, 13, 14, 60, 10, 10, 10, 10, 10, 10, 10, 10, 10, 12, 14, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(2, i).column_letter].width = w
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "F3"
    return last


def _summary(wb, last):
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Summary (computed from Results)"
    sm["A1"].font = _font(bold=True, size=14)
    for c, h in enumerate(["Metric", "Fable (reference)", "Qwen base", "Qwen tuned", "Change (tuned - base)"], 1):
        cell = sm.cell(3, c, h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = _hdr
        cell.alignment = _center

    def rng(col):
        return f"Results!${col}${FIRST_ROW}:${col}${last}"

    metrics = [
        ("Prompts scored", lambda ce, re, fi: f"=COUNT({rng(ce)})", "0"),
        ("Total compile errors", lambda ce, re, fi: f"=SUM({rng(ce)})", "0"),
        ("Total runtime errors", lambda ce, re, fi: f"=SUM({rng(re)})", "0"),
        ("Total errors", lambda ce, re, fi: f"=SUM({rng(ce)})+SUM({rng(re)})", "0"),
        ("Errors per prompt", lambda ce, re, fi: f'=IFERROR((SUM({rng(ce)})+SUM({rng(re)}))/COUNT({rng(ce)}),"-")', "0.00"),
        ("Clean runs (0 compile, 0 runtime) %", lambda ce, re, fi: f'=IFERROR(COUNTIFS({rng(ce)},0,{rng(re)},0)/COUNT({rng(ce)}),"-")', "0.0%"),
        ("Average fidelity (1-5)", lambda ce, re, fi: f'=IFERROR(AVERAGE({rng(fi)}),"-")', "0.00"),
    ]
    for i, (label, fn, fmt) in enumerate(metrics):
        r = 4 + i
        sm.cell(r, 1, label).font = _font()
        for j, m in enumerate(MODELS):
            cell = sm.cell(r, 2 + j, fn(*MODEL_COLS[m]))
            cell.font = _font()
            cell.number_format = fmt
            cell.alignment = _center
        cell = sm.cell(r, 5, f'=IF($D$4=0,"-",IFERROR(D{r}-C{r},"-"))')
        cell.font = _font()
        cell.number_format = fmt
        cell.alignment = _center

    r = 4 + len(metrics) + 1
    fl, sc, qb = HUMAN_FLAG_COL, HUMAN_SCORE_COL, MODEL_COLS["qwen_base"][2]
    sm.cell(r, 1, "Judge check (Fable vs human, Qwen base)").font = _font(bold=True)
    sm.cell(r + 1, 1, "Rows human-scored").font = _font()
    sm.cell(r + 1, 2, f'=SUMPRODUCT(({rng(fl)}="Y")*ISNUMBER({rng(sc)}))').font = _font()
    sm.cell(r + 2, 1, "Mean |Fable - human| fidelity gap").font = _font()
    c = sm.cell(r + 2, 2, f'=IFERROR(SUMPRODUCT(({rng(fl)}="Y")*ISNUMBER({rng(sc)})*ABS({rng(qb)}-{rng(sc)}))/B{r + 1},"-")')
    c.font, c.number_format = _font(), "0.00"
    sm.cell(r + 3, 1, "Share where Fable and human agree exactly").font = _font()
    c = sm.cell(r + 3, 2, f'=IFERROR(SUMPRODUCT(({rng(fl)}="Y")*ISNUMBER({rng(sc)})*({rng(qb)}={rng(sc)}))/B{r + 1},"-")')
    c.font, c.number_format = _font(), "0.0%"
    sm.cell(r + 4, 1, "A gap under about 0.5 suggests Fable's fidelity scores can be trusted; larger gaps mean the rubric or judge prompt needs work.").font = _font(italic=True, color="7F7F7F")

    r = r + 6
    sm.cell(r, 1, "Errors by template (Qwen base vs tuned)").font = _font(bold=True)
    for c, h in enumerate(["Template", "Qwen base errors", "Qwen tuned errors", "Qwen base avg fidelity", "Qwen tuned avg fidelity"], 1):
        cell = sm.cell(r + 1, c, h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = _hdr
        cell.alignment = _center
    n_templates = (last - FIRST_ROW + 1) // 10
    for t in range(n_templates):
        rr = r + 2 + t
        a, b = FIRST_ROW + t * 10, FIRST_ROW + t * 10 + 9
        sm.cell(rr, 1, f"=Results!A{a}").font = _font()
        for c, (m, kind) in enumerate([("qwen_base", "err"), ("qwen_tuned", "err"), ("qwen_base", "fid"), ("qwen_tuned", "fid")], 2):
            ce, re, fi = MODEL_COLS[m]
            if kind == "err":
                f = f"=SUM(Results!{ce}{a}:{ce}{b})+SUM(Results!{re}{a}:{re}{b})"
                fmt = "0"
            else:
                f = f'=IFERROR(AVERAGE(Results!{fi}{a}:{fi}{b}),"-")'
                fmt = "0.00"
            cell = sm.cell(rr, c, f)
            cell.font, cell.number_format, cell.alignment = _font(), fmt, _center
    r = r + 2 + n_templates + 1
    for title, col, values in [
        ("Errors by customer segment (Qwen base vs tuned)", SEGMENT_COL, ["individual", "startup", "institutional", "reference"]),
        ("Errors by prompt specificity (Qwen base vs tuned)", SPEC_COL, ["vague", "precise", "over_specified", "ceiling"]),
    ]:
        sm.cell(r, 1, title).font = _font(bold=True)
        for c, h in enumerate(["Group", "Qwen base errors", "Qwen tuned errors", "Qwen base avg fidelity", "Qwen tuned avg fidelity"], 1):
            cell = sm.cell(r + 1, c, h)
            cell.font = _font(bold=True, color="FFFFFF")
            cell.fill = _hdr
            cell.alignment = _center
        for k, v in enumerate(values):
            rr = r + 2 + k
            sm.cell(rr, 1, v).font = _font()
            for c, (m, kind) in enumerate([("qwen_base", "err"), ("qwen_tuned", "err"), ("qwen_base", "fid"), ("qwen_tuned", "fid")], 2):
                ce, re, fi = MODEL_COLS[m]
                if kind == "err":
                    f = f'=SUMIFS({rng(ce)},{rng(col)},"{v}")+SUMIFS({rng(re)},{rng(col)},"{v}")'
                    fmt = "0"
                else:
                    f = f'=IFERROR(AVERAGEIFS({rng(fi)},{rng(col)},"{v}"),"-")'
                    fmt = "0.00"
                cell = sm.cell(rr, c, f)
                cell.font, cell.number_format, cell.alignment = _font(), fmt, _center
        r = r + 2 + len(values) + 1
    sm.column_dimensions["A"].width = 44
    for col in "BCDE":
        sm.column_dimensions[col].width = 22


def build(prompts, path):
    wb = Workbook()
    _legend(wb)
    last = _results(wb, prompts)
    _summary(wb, last)
    wb.save(path)


def fill(path, prompts, results, human_rows=None):
    """results: {model: {prompt_id: {"compile_errors", "runtime_errors", "fidelity", "note"}}}"""
    wb = load_workbook(path)
    ws = wb["Results"]
    row_of = {p["id"]: FIRST_ROW + i for i, p in enumerate(prompts)}
    for model, by_id in results.items():
        ce, re, fi = MODEL_COLS[model]
        for pid, res in by_id.items():
            r = row_of[pid]
            ws[f"{ce}{r}"] = res.get("compile_errors")
            ws[f"{re}{r}"] = res.get("runtime_errors")
            ws[f"{fi}{r}"] = res.get("fidelity")
            if model == "qwen_base" and res.get("note"):
                ws[f"{NOTES_COL}{r}"] = res["note"][:300]
    for pid in human_rows or []:
        r = row_of[pid]
        if ws[f"{HUMAN_FLAG_COL}{r}"].value is None:
            ws[f"{HUMAN_FLAG_COL}{r}"] = "Y"
    wb.save(path)
